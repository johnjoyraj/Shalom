"""
Shalom Kerala Stores - Inventory & Sales Excel Sync Engine
Synchronizes online orders into SHALOM.xlsx (Sales_Invoicing & Inventory_Master)
"""

import openpyxl
import json
import os
from datetime import datetime
from build_catalog import build_catalog

EXCEL_FILE = "SHALOM.xlsx"

def sync_order_to_excel(order_data, excel_path=EXCEL_FILE):
    """
    Takes an order dictionary from the web store:
    1. Appends each item to 'Sales_Invoicing' sheet
    2. Decrements the item's 'Current stock' in 'Inventory_Master' sheet
    3. Recomputes Stock Valuation
    """
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found.")
        return False

    print(f"Opening {excel_path} for inventory synchronization...")
    wb = openpyxl.load_workbook(excel_path)
    
    ws_inv = wb['Inventory_Master']
    ws_sales = wb['Sales_Invoicing']

    # Build item name / code map to row index in Inventory_Master
    # Col 1: Item Name, Col 4: Item code, Col 12: Current stock, Col 9: Purchase Price, Col 16: Stock Valuation
    item_row_map = {}
    for r_idx in range(2, ws_inv.max_row + 1):
        name = ws_inv.cell(row=r_idx, column=1).value
        code = ws_inv.cell(row=r_idx, column=4).value
        if name:
            item_row_map[str(name).strip().upper()] = r_idx
        if code:
            item_row_map[str(code).strip().upper()] = r_idx

    invoice_no = order_data.get('invoiceNumber', f"INV-{datetime.now().strftime('%Y%m%d-%H%M%S')}")
    order_date = datetime.now()
    customer_name = order_data.get('customer', {}).get('name', 'Walk-in / Online Customer')
    
    items = order_data.get('items', [])
    updated_items_count = 0

    for item in items:
        item_name = item.get('name', '').strip()
        item_code = str(item.get('code', '')).strip()
        qty = float(item.get('quantity', 1))
        rate = float(item.get('price', 0))
        gst_rate = float(item.get('gstRate', 0))
        taxable_val = round(rate * qty, 2)
        cgst = round((taxable_val * (gst_rate / 2)) / 100, 2) if gst_rate > 0 else 0
        sgst = cgst
        total_amt = round(taxable_val + cgst + sgst, 2)

        # 1. Append row to Sales_Invoicing
        # Headers: ['Invoice No.', 'Date', 'Customer Name', 'Item Name', 'HSN Code', 'Quantity', 'Rate (₹)', 'Taxable Value (₹)', 'GST Rate (%)', 'CGST (₹)', 'SGST (₹)', 'Total Bill Amount (₹)']
        sales_next_row = ws_sales.max_row + 1
        ws_sales.cell(row=sales_next_row, column=1, value=invoice_no)
        ws_sales.cell(row=sales_next_row, column=2, value=order_date)
        ws_sales.cell(row=sales_next_row, column=3, value=customer_name)
        ws_sales.cell(row=sales_next_row, column=4, value=item_name)
        ws_sales.cell(row=sales_next_row, column=5, value=item.get('hsn', ''))
        ws_sales.cell(row=sales_next_row, column=6, value=qty)
        ws_sales.cell(row=sales_next_row, column=7, value=rate)
        ws_sales.cell(row=sales_next_row, column=8, value=taxable_val)
        ws_sales.cell(row=sales_next_row, column=9, value=gst_rate)
        ws_sales.cell(row=sales_next_row, column=10, value=cgst)
        ws_sales.cell(row=sales_next_row, column=11, value=sgst)
        ws_sales.cell(row=sales_next_row, column=12, value=total_amt)

        # 2. Decrement stock in Inventory_Master
        lookup_key = item_name.upper()
        target_row = item_row_map.get(lookup_key) or item_row_map.get(item_code.upper())
        if target_row:
            current_stock_cell = ws_inv.cell(row=target_row, column=12)
            try:
                curr_stock = float(current_stock_cell.value or 0)
            except (ValueError, TypeError):
                curr_stock = 0.0
            
            new_stock = max(0.0, curr_stock - qty)
            current_stock_cell.value = new_stock

            # Update Stock Valuation (col 16) = Current stock * Purchase price (col 9)
            purchase_price = ws_inv.cell(row=target_row, column=9).value or 0
            try:
                val = new_stock * float(purchase_price)
                ws_inv.cell(row=target_row, column=16, value=round(val, 2))
            except (ValueError, TypeError):
                pass

            # Update Stock Status Alert (col 17)
            alert_cell = ws_inv.cell(row=target_row, column=17)
            low_limit = ws_inv.cell(row=target_row, column=13).value or 1
            try:
                low_limit_val = float(low_limit)
            except (ValueError, TypeError):
                low_limit_val = 1.0

            if new_stock <= 0:
                alert_cell.value = "Out of Stock"
            elif new_stock <= low_limit_val:
                alert_cell.value = "Low Stock"
            else:
                alert_cell.value = "In Stock"

            updated_items_count += 1
            print(f"  Decremented stock for '{item_name}': {curr_stock} -> {new_stock}")
        else:
            print(f"  Warning: Item '{item_name}' (SKU: {item_code}) not found in Inventory_Master.")

    wb.save(excel_path)
    print(f"Successfully saved {excel_path} with {updated_items_count} inventory updates.")

    # Re-run build_catalog to refresh web products
    print("Re-building web store catalog...")
    build_catalog()
    return True

if __name__ == '__main__':
    # Test sample synchronization
    sample_order = {
        "invoiceNumber": f"INV-TEST-{datetime.now().strftime('%Y%m%d%H%M')}",
        "customer": {"name": "Test Customer (Gorakhpur)"},
        "items": [
            {
                "name": "120W TYPE C CABLE",
                "code": "514182562487",
                "quantity": 1,
                "price": 60.0,
                "gstRate": 0.0
            }
        ]
    }
    print("Running inventory sync engine...")
    sync_order_to_excel(sample_order)
